from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
import os
import numpy as np
import pandas as pd

from .measure_ops import (
    OP_COLUMNS,
    OP_DIMENSION_COL,
    OP_FINGER_COL,
    OP_MARKER_1_COL,
    OP_MARKER_2_COL,
    OP_TYPE_COL,
    drop_invalid_marker_ops,
    is_invalid_ops_marker,
    normalize_ops_columns,
)
from .paths import asset_path
from .output_cleanup import detections_labeled_csv


@dataclass
class HandMeasuresResult:
    ok: bool
    message: str
    detailed_xlsx: Path | None = None
    n_ok: int = 0
    n_missing: int = 0
    n_invalid: int = 0
    missing_markers: list[str] | None = None


def _safe_unlink(path: Path) -> bool:
    try:
        if path.exists() and path.is_file():
            path.unlink()
            return True
    except Exception:
        return False
    return False


def _safe_to_styled_detailed_xlsx(df: pd.DataFrame, path: Path) -> Path:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font, PatternFill
        from openpyxl.utils import get_column_letter
    except ImportError as e:
        raise RuntimeError("openpyxl is required to generate hand_measures_detailed.xlsx") from e

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = "Detailed measures"

    headers = list(df.columns)
    ws.append(headers)

    for _, row in df.iterrows():
        values = []
        for value in row.tolist():
            if pd.isna(value):
                values.append(None)
            else:
                values.append(value)
        ws.append(values)

    header_fill = PatternFill("solid", fgColor="E8EDF3")
    header_font = Font(bold=True, color="1F2937")
    distance_fill = PatternFill("solid", fgColor="DFF7E5")
    distance_header_fill = PatternFill("solid", fgColor="C8F2D0")
    distance_font = Font(bold=True, color="14532D")

    distance_col = headers.index("DIST_EUCLIDEA") + 1 if "DIST_EUCLIDEA" in headers else None

    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="left", vertical="center")

    if distance_col is not None:
        header_cell = ws.cell(row=1, column=distance_col)
        header_cell.fill = distance_header_fill
        header_cell.font = distance_font

        for row_idx in range(2, ws.max_row + 1):
            cell = ws.cell(row=row_idx, column=distance_col)
            cell.fill = distance_fill
            cell.font = Font(color="14532D")
            cell.alignment = Alignment(horizontal="right", vertical="center")
            cell.number_format = "0.000"

    for col_idx, header in enumerate(headers, start=1):
        letter = get_column_letter(col_idx)
        values = [header]
        for row_idx in range(2, ws.max_row + 1):
            value = ws.cell(row=row_idx, column=col_idx).value
            values.append("" if value is None else str(value))
        width = max(len(str(v)) for v in values) + 2
        if col_idx == distance_col:
            width = max(width, 18)
        ws.column_dimensions[letter].width = min(width, 42)

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions

    tmp = path.with_name(f"{path.stem}.tmp{path.suffix}")
    try:
        wb.save(tmp)
        try:
            os.replace(tmp, path)
            return path
        except PermissionError:
            stamped = path.with_name(f"{path.stem}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{path.suffix}")
            os.replace(tmp, stamped)
            return stamped
    finally:
        if tmp.exists():
            try:
                tmp.unlink()
            except Exception:
                pass


def _read_ops_csv(path: Path) -> pd.DataFrame:
    last_error = None
    for encoding in ("utf-8-sig", "utf-8", "latin1"):
        try:
            return pd.read_csv(path, sep=";", encoding=encoding)
        except Exception as e:
            last_error = e
    raise last_error


def _is_invalid_marker(m: str | None) -> bool:
    return is_invalid_ops_marker(m)


def compute_hand_measures(
    out_dir: Path,
    labeled_csv: Path | None = None,
    ops_csv: Path | None = None,
    log=None,
) -> HandMeasuresResult:
    try:
        out_dir = Path(out_dir)
        out_dir.mkdir(parents=True, exist_ok=True)

        if labeled_csv:
            labeled_csv = Path(labeled_csv)
        else:
            labeled_csv = detections_labeled_csv(out_dir, create=False)
            if not labeled_csv.exists():
                labeled_csv = out_dir / "detections_labeled.csv"
        if not labeled_csv.exists():
            return HandMeasuresResult(False, f"Not found: {labeled_csv}")

        ops_csv = Path(ops_csv) if ops_csv else asset_path("data", "LandmarksOperations_data.csv")
        if not ops_csv.exists():
            return HandMeasuresResult(False, f"Not found: {ops_csv}\n(Place it in src/assets/data/LandmarksOperations_data.csv)")

        if log:
            log(f"[INFO] detections_labeled: {labeled_csv}")
            log(f"[INFO] operations:         {ops_csv}")

        # ----------------------------
        # 1) load detections
        # ----------------------------
        det = pd.read_csv(labeled_csv)

        label_col = None
        for c in ("label", "Label", "LABEL", "name", "Name"):
            if c in det.columns:
                label_col = c
                break
        if label_col is None:
            return HandMeasuresResult(False, f"No encuentro columna de etiqueta en {labeled_csv}. Columnas: {list(det.columns)}")

        xyz_cols = None
        for triple in (("x", "y", "z"), ("X", "Y", "Z")):
            if all(c in det.columns for c in triple):
                xyz_cols = triple
                break
        if xyz_cols is None:
            return HandMeasuresResult(False, f"No encuentro columnas XYZ en {labeled_csv}. Columnas: {list(det.columns)}")

        det[label_col] = det[label_col].astype(str).str.strip()
        det = det[det[label_col].astype(bool)].copy()

        # if there are duplicates, keep the first
        det = det.drop_duplicates(subset=[label_col], keep="first")

        label_to_point = {
            row[label_col]: np.array([row[xyz_cols[0]], row[xyz_cols[1]], row[xyz_cols[2]]], dtype=float)
            for _, row in det.iterrows()
        }

        def euclidean(p1: np.ndarray, p2: np.ndarray) -> float:
            return float(np.linalg.norm(p1 - p2))

        if log:
            log(f"[INFO] points loaded: {len(label_to_point)} (label_col='{label_col}', xyz={xyz_cols})")

        # ----------------------------
        # 2) load recipes
        # ----------------------------
        # The operations CSV uses ';'. Prefer UTF-8 so accented labels are not
        # accidentally mojibake-decoded by latin1, which accepts almost anything.
        ops = _read_ops_csv(ops_csv)

        ops, missing = normalize_ops_columns(ops)
        if missing:
            return HandMeasuresResult(False, f"Missing columns {missing} in {ops_csv}. Columns: {list(ops.columns)}")

        ops = ops[OP_COLUMNS].copy()
        ops[OP_TYPE_COL] = ops[OP_TYPE_COL].ffill()
        ops[OP_FINGER_COL] = ops[OP_FINGER_COL].ffill()

        def norm_cell(s):
            if pd.isna(s):
                return None
            s = str(s).strip()
            return s if s else None

        for c in OP_COLUMNS:
            ops[c] = ops[c].apply(norm_cell)

        ops = ops.dropna(subset=[OP_DIMENSION_COL]).copy()
        ops, skipped_invalid_ops = drop_invalid_marker_ops(ops)
        if skipped_invalid_ops and log:
            log(f"[INFO] skipped unsupported line-based operations: {skipped_invalid_ops}")

        # ----------------------------
        # 3) compute distances
        # ----------------------------
        rows = []
        missing_markers = set()
        invalid_recipes = 0

        for _, r in ops.iterrows():
            tipo = r[OP_TYPE_COL]
            finger = r[OP_FINGER_COL]
            dim = r[OP_DIMENSION_COL]
            m1 = r[OP_MARKER_1_COL]
            m2 = r[OP_MARKER_2_COL]

            if _is_invalid_marker(m1) or _is_invalid_marker(m2):
                rows.append({
                    OP_TYPE_COL: tipo, OP_FINGER_COL: finger, OP_DIMENSION_COL: dim,
                    OP_MARKER_1_COL: m1, OP_MARKER_2_COL: m2,
                    "DIST_EUCLIDEA": np.nan,
                    "STATUS": "INVALID_RECIPE",
                })
                invalid_recipes += 1
                continue

            p1 = label_to_point.get(m1)
            p2 = label_to_point.get(m2)

            if p1 is None or p2 is None:
                dist = np.nan
                status = "MISSING_MARKER"
                if p1 is None:
                    missing_markers.add(m1)
                if p2 is None:
                    missing_markers.add(m2)
            else:
                dist = euclidean(p1, p2)
                status = "OK"

            rows.append({
                OP_TYPE_COL: tipo, OP_FINGER_COL: finger, OP_DIMENSION_COL: dim,
                OP_MARKER_1_COL: m1, OP_MARKER_2_COL: m2,
                "DIST_EUCLIDEA": dist,
                "STATUS": status,
            })

        measures = pd.DataFrame(rows)

        out_detailed_xlsx = out_dir / "hand_measures_detailed.xlsx"

        # Older versions generated this file. The XLSX is now the detailed
        # user-facing result, and the wide CSV has been removed because its
        # format was confusing. Drop stale copies from previous runs.
        for stale in (out_dir / "hand_measures_detailed.csv", out_dir / "hand_measures_wide.csv"):
            if stale.exists() and not _safe_unlink(stale) and log:
                log(f"[WARN] Could not remove old result file: {stale}")

        saved_detailed_xlsx = _safe_to_styled_detailed_xlsx(measures, out_detailed_xlsx)

        n_ok = int((measures["STATUS"] == "OK").sum())
        n_miss = int((measures["STATUS"] == "MISSING_MARKER").sum())
        n_inv = int((measures["STATUS"] == "INVALID_RECIPE").sum())

        if log:
            log(f"[OK] detailed: {saved_detailed_xlsx}")
            log(f"[SUMMARY] OK={n_ok} | MISSING={n_miss} | INVALID={n_inv}")

        return HandMeasuresResult(
            ok=True,
            message="OK",
            detailed_xlsx=saved_detailed_xlsx,
            n_ok=n_ok,
            n_missing=n_miss,
            n_invalid=n_inv,
            missing_markers=sorted(missing_markers),
        )

    except Exception as e:
        return HandMeasuresResult(False, str(e))
