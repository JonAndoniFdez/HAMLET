from __future__ import annotations

from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUT_DIR = ROOT / "docs"
DOCX_PATH = OUT_DIR / "Automatic_Hand_Measurement_Tool_User_Guide.docx"


BLUE = RGBColor(46, 116, 181)
DARK_BLUE = RGBColor(31, 77, 120)
INK = RGBColor(32, 32, 32)
MUTED = RGBColor(90, 90, 90)
LIGHT_BLUE_FILL = "E8EEF5"
LIGHT_GRAY_FILL = "F2F4F7"
CALLOUT_FILL = "F4F6F9"
BORDER = "B7C9DD"


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table, color="B7C9DD", size="6") -> None:
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_table_width(table, widths_in: list[float]) -> None:
    table.autofit = False
    for row in table.rows:
        for idx, width in enumerate(widths_in):
            if idx >= len(row.cells):
                continue
            row.cells[idx].width = Inches(width)


def set_repeat_table_header(row) -> None:
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def style_table(table, widths: list[float], header_fill=LIGHT_BLUE_FILL) -> None:
    set_table_borders(table, BORDER)
    set_table_width(table, widths)
    if table.rows:
        set_repeat_table_header(table.rows[0])
        for cell in table.rows[0].cells:
            set_cell_shading(cell, header_fill)
            for p in cell.paragraphs:
                for run in p.runs:
                    run.bold = True
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.1


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[float]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for i, header in enumerate(headers):
        table.rows[0].cells[i].text = header
    for row in rows:
        cells = table.add_row().cells
        for i, value in enumerate(row):
            cells[i].text = str(value)
    style_table(table, widths)


def add_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    return p


def add_body(doc: Document, text: str):
    p = doc.add_paragraph(text)
    p.style = doc.styles["Normal"]
    return p


def add_bullets(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(item, style="List Bullet")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25


def add_numbers(doc: Document, items: list[str]) -> None:
    for item in items:
        p = doc.add_paragraph(item, style="List Number")
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.25


def add_callout(doc: Document, title: str, text: str, fill: str = CALLOUT_FILL) -> None:
    table = doc.add_table(rows=1, cols=1)
    table.autofit = False
    table.rows[0].cells[0].width = Inches(6.5)
    set_table_borders(table, "D5DCE6", "4")
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, fill)
    set_cell_margins(cell, top=120, bottom=120, start=160, end=160)
    p = cell.paragraphs[0]
    r = p.add_run(title)
    r.bold = True
    r.font.color.rgb = DARK_BLUE
    r.font.size = Pt(10.5)
    p.add_run(f" {text}")
    p.paragraph_format.space_after = Pt(0)


def configure_styles(doc: Document) -> None:
    section = doc.sections[0]
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    normal.font.size = Pt(11)
    normal.font.color.rgb = INK
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.25

    for name, size, color, before, after in [
        ("Heading 1", 16, BLUE, 18, 10),
        ("Heading 2", 13, BLUE, 14, 7),
        ("Heading 3", 12, DARK_BLUE, 10, 5),
    ]:
        style = doc.styles[name]
        style.font.name = "Calibri"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True

    for name in ("List Bullet", "List Number"):
        style = doc.styles[name]
        style.font.name = "Calibri"
        style.font.size = Pt(11)
        style.paragraph_format.left_indent = Inches(0.375)
        style.paragraph_format.first_line_indent = Inches(-0.188)
        style.paragraph_format.space_after = Pt(4)
        style.paragraph_format.line_spacing = 1.25


def add_footer(section, label: str) -> None:
    footer = section.footer
    p = footer.paragraphs[0]
    p.text = label
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p.runs[0].font.size = Pt(9)
    p.runs[0].font.color.rgb = MUTED


def add_cover(doc: Document) -> None:
    logo = ROOT / "src" / "assets" / "icons" / "uji_logo.png"
    if logo.exists():
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        try:
            p.add_run().add_picture(str(logo), width=Inches(1.65))
        except Exception:
            pass

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(92)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run("HAMLET: Hand Anthropometric Measurement and Landmark Extraction Tool")
    run.bold = True
    run.font.size = Pt(26)
    run.font.color.rgb = DARK_BLUE

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(24)
    run = p.add_run("Guía de usuario completa")
    run.font.size = Pt(16)
    run.font.color.rgb = BLUE

    add_callout(
        doc,
        "Propósito de la guía.",
        "Este documento explica de forma completa cómo instalar, iniciar y utilizar la aplicación en su versión actual. "
        "Está pensado para que una persona que no conozca el proyecto pueda entender el flujo completo de trabajo, "
        "los archivos de entrada, los resultados generados y los puntos de revisión manual necesarios para obtener medidas fiables.",
    )

    metadata = [
        ["Aplicación", "HAMLET: Hand Anthropometric Measurement and Landmark Extraction Tool"],
        ["Tipo de documento", "Guía de usuario"],
        ["Fecha de emisión", date.today().strftime("%Y-%m-%d")],
        ["Responsable", "Jon Andoni Fernández Moyano, Universitat Jaume I (UJI)"],
        ["Ámbito", "Uso de la aplicación ejecutable en Windows"],
    ]
    add_heading(doc, "Datos del documento", 2)
    add_table(doc, ["Campo", "Valor"], metadata, [1.8, 4.7])

    doc.add_page_break()


def add_index(doc: Document) -> None:
    add_heading(doc, "Índice de contenidos", 1)
    sections = [
        "1. Alcance y visión general",
        "2. Conceptos básicos y criterios de uso",
        "3. Requisitos e instalación",
        "4. Entrada, salida y estructura de carpetas",
        "5. Guía rápida",
        "6. Guía extendida por pantallas",
        "7. Reabrir un proyecto existente",
        "8. Formato de CSV de medidas",
        "9. FAQ y resolución de problemas",
        "10. Contacto",
    ]
    add_bullets(doc, sections)
    add_callout(
        doc,
        "Nota para edición.",
        "El índice es manual para que el archivo sea estable al abrirlo en distintos equipos. "
        "Si se desea, puede sustituirse por una tabla de contenido automática desde Microsoft Word.",
    )
    doc.add_page_break()


def add_summary(doc: Document) -> None:
    add_heading(doc, "1. Alcance y visión general", 1)
    add_body(
        doc,
        "HAMLET: Hand Anthropometric Measurement and Landmark Extraction Tool es una aplicación de escritorio para Windows orientada a obtener "
        "medidas antropométricas de una mano a partir de un modelo 3D en formato OBJ y su textura asociada. "
        "La herramienta combina detección automática, mapeo 2D a 3D, validación manual, rigging asistido por "
        "OpenPose, etiquetado automático, corrección interactiva y cálculo final de distancias.",
    )
    add_body(
        doc,
        "La guía cubre el uso de la versión ejecutable de la aplicación: instalación, selección de archivos, "
        "detección y revisión de marcadores, rigging, etiquetado, selección de medidas, cálculo final y reapertura "
        "de proyectos ya procesados. No presupone conocimiento previo de versiones antiguas ni de documentación externa.",
    )
    add_heading(doc, "Flujo de trabajo de la aplicación", 2)
    add_numbers(
        doc,
        [
            "Welcome: iniciar una medición nueva o abrir una carpeta de salida de un proyecto anterior.",
            "Input Data: seleccionar el modelo OBJ, la textura y la carpeta de salida.",
            "Color Threshold Selection: elegir el preset azul mate o generar umbrales HSV mediante selección manual.",
            "Landmark Detection + 2D/3D Validation: detectar marcadores, mapearlos sobre la malla y validarlos en 2D y 3D.",
            "Rigging and Labeling: generar el rig, ejecutar OpenPose, etiquetar automáticamente y corregir etiquetas manualmente.",
            "Measures selector: seleccionar las medidas que se calcularán y, si procede, añadir medidas personalizadas.",
            "Calculation of measures: calcular distancias, revisar una vista previa y abrir la carpeta de resultados.",
        ],
    )
    add_callout(
        doc,
        "Idea clave.",
        "La aplicación automatiza gran parte del proceso, pero la calidad final depende de revisar los landmarks "
        "y las etiquetas antes de calcular las medidas. La validación visual no es opcional cuando se buscan resultados fiables.",
    )


def add_concepts(doc: Document) -> None:
    add_heading(doc, "2. Conceptos básicos y criterios de uso", 1)
    add_body(
        doc,
        "Antes de iniciar una medición conviene distinguir los elementos que intervienen en el flujo. "
        "La aplicación trabaja con una textura 2D, una malla 3D y una lista final de landmarks etiquetados. "
        "Cada medida se calcula como una distancia entre dos etiquetas concretas.",
    )
    rows = [
        ["Modelo OBJ", "Malla 3D de la mano. Debe contener geometría válida y coordenadas UV."],
        ["Textura", "Imagen asociada al modelo. Sobre ella se detectan inicialmente los marcadores de color."],
        ["Landmark", "Punto anatómico detectado o añadido por el usuario. Puede existir en 2D, en 3D o en ambos."],
        ["Mapeo 2D/3D", "Proceso que proyecta los puntos detectados en la textura sobre la superficie del modelo."],
        ["Rigging", "Construcción de una estructura de apoyo para orientar y etiquetar la mano."],
        ["Etiqueta", "Nombre anatómico asignado a un landmark. Las medidas dependen de estas etiquetas."],
        ["Receta de medida", "Fila del CSV de operaciones que indica qué dos etiquetas se usan para calcular una distancia."],
    ]
    add_table(doc, ["Concepto", "Significado práctico"], rows, [1.8, 4.7])
    add_heading(doc, "Responsabilidades del usuario", 2)
    add_bullets(
        doc,
        [
            "Comprobar que el OBJ, el MTL y la textura pertenecen al mismo modelo y están bien vinculados.",
            "Verificar visualmente que los puntos detectados corresponden a marcadores reales.",
            "Corregir puntos o etiquetas cuando la detección automática no sea suficiente.",
            "Revisar el CSV final antes de usar las medidas en análisis externos.",
        ],
    )
    add_heading(doc, "Cuándo repetir un paso", 2)
    add_table(
        doc,
        ["Situación", "Acción recomendada"],
        [
            ["Muchos falsos positivos en la textura", "Volver al umbral y usar un preset más estricto o selección manual más precisa."],
            ["Puntos correctos en 2D pero mal ubicados en 3D", "Revisar la pestaña 3D y añadir/eliminar puntos directamente sobre la malla."],
            ["OpenPose muestra una estructura incoherente", "Revisar orientación del modelo y repetir Automatic Rigging."],
            ["Muchas etiquetas incorrectas", "Usar Interactive Label Editor y, si el problema es general, revisar orientación y rigging."],
            ["Medidas con MISSING_MARKER", "Comprobar que las etiquetas requeridas existen en detections_labeled.csv."],
        ],
        [2.3, 4.2],
    )


def add_requirements(doc: Document) -> None:
    add_heading(doc, "3. Requisitos e instalación", 1)
    add_heading(doc, "Requisitos del sistema", 2)
    rows = [
        ["Sistema operativo", "Windows 10 o Windows 11 de 64 bits."],
        ["CPU", "Procesador de 64 bits. El rendimiento depende del tamaño de la malla y de la textura."],
        ["Memoria RAM", "Mínimo recomendado: 8 GB. Recomendado para uso fluido: 16 GB o más."],
        ["Almacenamiento", "La release portable puede ocupar varios GB si incluye Blender y los modelos de OpenPose. Reservar espacio adicional para resultados."],
        ["Permisos", "Usar una carpeta con permisos de lectura/escritura. Evitar ubicaciones bloqueadas por políticas corporativas."],
    ]
    add_table(doc, ["Requisito", "Descripción"], rows, [1.7, 4.8])
    add_heading(doc, "Instalación de la versión ejecutable", 2)
    add_numbers(
        doc,
        [
            "Descargar el archivo ZIP publicado en la release.",
            "Descomprimir el ZIP en una carpeta local, por ejemplo Documentos o una carpeta de trabajo del usuario.",
            "Abrir la carpeta descomprimida completa.",
            "Ejecutar HAMLET.exe.",
        ],
    )
    add_callout(
        doc,
        "No copiar sólo el .exe.",
        "La aplicación portable necesita las carpetas que acompañan al ejecutable, como assets, blender, openpose, PySide6 y vtkmodules. "
        "Para mover la aplicación a otro equipo, copiar o distribuir la carpeta completa o el ZIP completo.",
    )
    add_heading(doc, "Estructura esperada de la carpeta portable", 2)
    rows = [
        ["HAMLET.exe", "Ejecutable principal de la aplicación."],
        ["assets/", "Iconos, scripts de Blender, plantillas CSV y visor 3D local."],
        ["blender/", "Blender portable usado por los pasos que necesitan generar renders, rigs o escenas .blend."],
        ["openpose/models/hand/", "Modelos Caffe de mano usados por OpenCV DNN para detectar keypoints."],
        ["Paquetes Python", "Bibliotecas incluidas por PyInstaller: PySide6, OpenCV, pandas, numpy, VTK/Vedo, etc."],
    ]
    add_table(doc, ["Elemento", "Uso"], rows, [2.0, 4.5])


def add_io(doc: Document) -> None:
    add_heading(doc, "4. Entrada, salida y estructura de carpetas", 1)
    add_heading(doc, "Datos de entrada", 2)
    add_bullets(
        doc,
        [
            "Modelo 3D de la mano en formato .OBJ.",
            "Archivo .MTL asociado cuando el OBJ lo utilice para referenciar materiales.",
            "Textura asociada en formato .PNG, .JPG, .JPEG, .TIF, .TIFF, .BMP o .WEBP.",
            "Carpeta de salida donde se guardará todo el proceso.",
        ],
    )
    add_callout(
        doc,
        "Buenas prácticas de entrada.",
        "Mantener OBJ, MTL y textura en la misma carpeta ayuda a evitar rutas rotas. "
        "La aplicación intenta detectar automáticamente una textura con el mismo nombre que el OBJ.",
    )
    add_heading(doc, "Resultados principales", 2)
    rows = [
        ["detections_all.csv", "Puntos detectados en 2D sobre la textura."],
        ["detections_mapped_3d.csv", "Puntos mapeados a coordenadas 3D sobre la malla."],
        ["hand_with_points.obj", "OBJ auxiliar con la mano y esferas de puntos usado por pasos posteriores."],
        ["detections_labeled.csv", "Puntos detectados con etiqueta final y coordenadas."],
        ["hand_rigged_labeled.blend", "Escena Blender final con rig y etiquetas aplicadas."],
        ["results/hand_measures_detailed.xlsx", "Resultado detallado con TYPE, FINGER, DIMENSION, marcadores, distancia y estado. La columna DIST_EUCLIDEA aparece resaltada."],
        ["results/hand_measures_wide.csv", "CSV en formato ancho, útil para análisis posterior."],
        ["results/points_xyz.csv", "Exportación simplificada de etiqueta y coordenadas XYZ."],
        ["project_context.json", "Metadatos para reabrir proyectos desde la pantalla Open project."],
    ]
    add_table(doc, ["Archivo o carpeta", "Contenido"], rows, [2.15, 4.35])
    add_heading(doc, "Estados habituales en medidas", 2)
    rows = [
        ["OK", "La medida se ha calculado correctamente."],
        ["MISSING_MARKER", "Falta una de las etiquetas requeridas por la medida."],
        ["INVALID_RECIPE", "La definición de la medida contiene marcadores no válidos o una receta incompatible."],
    ]
    add_table(doc, ["Estado", "Significado"], rows, [1.8, 4.7])


def add_quick_guide(doc: Document) -> None:
    add_heading(doc, "5. Guía rápida", 1)
    add_numbers(
        doc,
        [
            "Ejecutar HAMLET.exe desde la carpeta descomprimida de la release.",
            "Pulsar New measurement para iniciar un análisis nuevo.",
            "Seleccionar el archivo OBJ, la textura y la carpeta de salida.",
            "Elegir Blue matte color preset si los marcadores siguen el protocolo de color azul mate. Usar Manual color selection si la textura requiere ajustar el umbral.",
            "Esperar a que se complete Landmark Detection + 2D/3D Validation. Revisar los puntos en la ventana combinada 2D/3D antes de continuar.",
            "Ejecutar Rigging and Labeling. Revisar el resultado de OpenPose si está disponible y corregir las etiquetas en el editor interactivo.",
            "Seleccionar las medidas deseadas, añadir medidas personalizadas si hace falta y continuar.",
            "Revisar la vista previa final, abrir el CSV si se desea y pulsar Finish para abrir la carpeta de resultados.",
        ],
    )
    add_callout(
        doc,
        "Validación recomendada.",
        "Aunque el flujo es semiautomático, la revisión 2D/3D y la corrección manual de etiquetas son pasos clave para obtener medidas fiables.",
    )


def add_extended_guide(doc: Document) -> None:
    add_heading(doc, "6. Guía extendida por pantallas", 1)
    add_heading(doc, "Welcome", 2)
    add_body(
        doc,
        "Pantalla inicial de la aplicación. Permite iniciar una medición nueva, abrir un proyecto existente, consultar una ayuda rápida o ver información de la aplicación.",
    )
    add_table(
        doc,
        ["Control", "Función"],
        [
            ["New measurement", "Inicia el flujo normal desde Input Data."],
            ["Open project", "Permite seleccionar la carpeta de salida de un proyecto anterior y reabrir partes disponibles."],
            ["First time here?", "Muestra una guía breve dentro de la aplicación."],
            ["About", "Muestra versión, desarrollador y descripción de la herramienta."],
        ],
        [2.0, 4.5],
    )
    add_heading(doc, "Input Data", 2)
    add_body(
        doc,
        "En esta pantalla se definen los archivos de entrada y la carpeta donde se generará el proyecto. "
        "El botón Next sólo se habilita cuando el OBJ, la textura y la carpeta de salida son válidos.",
    )
    add_table(
        doc,
        ["Campo", "Descripción"],
        [
            ["3D Model (.obj)", "Modelo geométrico de la mano. Debe contener geometría y coordenadas UV."],
            ["Texture", "Imagen asociada al modelo. Si tiene el mismo nombre que el OBJ, puede detectarse automáticamente."],
            ["Output folder", "Carpeta donde se escriben CSV, OBJ auxiliares, BLEND y resultados finales."],
        ],
        [1.8, 4.7],
    )
    add_heading(doc, "Color Threshold Selection", 2)
    add_body(
        doc,
        "Define cómo se identificarán los marcadores sobre la textura. El modo recomendado es Blue matte color preset cuando los marcadores siguen el color del protocolo. "
        "Manual color selection abre una vista de textura donde el usuario selecciona muestras de color y genera rangos HSV personalizados.",
    )
    add_bullets(
        doc,
        [
            "High: umbral más restrictivo, útil cuando hay falsos positivos.",
            "Medium: opción equilibrada y recomendada para la mayoría de casos manuales.",
            "Low: tolerancia amplia, útil cuando hay variaciones de iluminación o color.",
            "Custom: permite editar H_TOL, S_TOL, V_TOL y PICK_RADIUS.",
        ],
    )
    add_callout(
        doc,
        "Selección manual.",
        "Hacer varios clics en zonas limpias y homogéneas de los marcadores. Evitar bordes, brillos y zonas mezcladas con color de piel o sombras.",
    )
    add_heading(doc, "Landmark Detection + 2D/3D Validation", 2)
    add_body(
        doc,
        "La aplicación detecta marcadores en la textura, los proyecta a la malla 3D y abre una validación combinada. "
        "La pantalla resume dos fases: Automatic Landmark Detection y 2D/3D Validation.",
    )
    add_table(
        doc,
        ["Elemento", "Uso"],
        [
            ["Automatic Landmark Detection", "Muestra los puntos iniciales 2D, el estado del mapeo a 3D y el total mapeado."],
            ["2D/3D Validation", "Permite revisar y editar el mismo conjunto de landmarks desde la textura y desde la malla."],
            ["Retry", "Repetir la detección automática o sólo la validación 2D/3D, según la información disponible."],
            ["Next", "Genera hand_with_points.obj y avanza al flujo de rigging y etiquetado."],
        ],
        [2.05, 4.45],
    )
    add_heading(doc, "Edición 2D", 3)
    add_bullets(
        doc,
        [
            "Útil para limpiar falsos positivos visibles sobre la textura.",
            "Permite añadir o eliminar puntos en la imagen 2D.",
            "Después de aplicar cambios 2D, la aplicación sincroniza el CSV 3D cuando es posible.",
        ],
    )
    add_heading(doc, "Edición 3D", 3)
    add_bullets(
        doc,
        [
            "Útil para confirmar que los puntos están realmente sobre la superficie de la mano.",
            "Shift + clic sobre la malla añade un nuevo punto en la superficie.",
            "Clic sobre un punto existente permite eliminarlo si es visible y seleccionable.",
            "Undo y Reset Camera ayudan a corregir errores de edición y orientación de vista.",
        ],
    )
    add_heading(doc, "Rigging and Labeling", 2)
    add_body(
        doc,
        "Esta pantalla agrupa Automatic Rigging, Automatic Labeling e Interactive Label Editor en un único flujo guiado. "
        "La pantalla muestra tarjetas de estado y permite reintentar fases concretas cuando existen los archivos intermedios necesarios.",
    )
    add_table(
        doc,
        ["Fase", "Qué hace"],
        [
            ["Automatic Rigging", "Renderiza una vista superior con Blender, ejecuta OpenPose sobre la imagen y construye el armature en Blender."],
            ["Automatic Labeling", "Asigna etiquetas a las esferas detectadas usando la plantilla de marcadores."],
            ["Interactive Label Editor", "Abre un editor 3D para revisar y corregir etiquetas manualmente."],
        ],
        [2.0, 4.5],
    )
    add_callout(
        doc,
        "Orientación del modelo.",
        "Si OpenPose o el etiquetado automático fallan de forma generalizada, revisar la orientación del modelo importado. "
        "El flujo asume una orientación compatible con el render superior usado por el rigging.",
    )
    add_heading(doc, "Interactive Label Editor", 3)
    add_bullets(
        doc,
        [
            "Clic izquierdo sobre un punto abre el diálogo Point editing.",
            "Nearby points muestra candidatos cercanos o más probables.",
            "All points permite buscar cualquier etiqueta disponible.",
            "Si se elige una etiqueta ya asignada a otro punto, la aplicación intercambia las etiquetas.",
            "ENTER o Continue finaliza el editor y guarda la corrección.",
        ],
    )
    add_heading(doc, "Measures selector", 2)
    add_body(
        doc,
        "Permite elegir qué medidas calcular. Por defecto carga LandmarksOperations_data.csv desde assets/data. "
        "Las medidas se agrupan por dedo y por tipo: Lengths, Breadths y Depths. También permite cargar un CSV personalizado.",
    )
    add_table(
        doc,
        ["Control", "Función"],
        [
            ["Browse", "Seleccionar un CSV de operaciones personalizado."],
            ["Select all", "Marca todas las medidas disponibles."],
            ["Select none", "Desmarca todas las medidas."],
            ["New Custom Measure", "Crea una medida extra entre dos etiquetas existentes."],
            ["Continue", "Guarda LandmarksOperations_selected.csv y avanza al cálculo."],
        ],
        [2.0, 4.5],
    )
    add_heading(doc, "Calculation of measures", 2)
    add_body(
        doc,
        "Calcula distancias euclídeas a partir de detections_labeled.csv y del CSV de operaciones seleccionado. "
        "La pantalla muestra una vista previa con las primeras medidas y un resumen de estados OK, MISSING e INVALID.",
    )
    add_bullets(
        doc,
        [
            "Open XLSX abre el resultado detallado generado.",
            "Finish copia o confirma los resultados finales, abre la carpeta results y cierra la aplicación.",
        ],
    )


def add_open_project(doc: Document) -> None:
    add_heading(doc, "7. Reabrir un proyecto existente", 1)
    add_body(
        doc,
        "La pantalla Open project permite seleccionar una carpeta de salida creada por una ejecución anterior. "
        "La aplicación inspecciona los archivos disponibles y habilita sólo las acciones que puede recuperar de forma segura.",
    )
    add_table(
        doc,
        ["Acción", "Requisitos principales"],
        [
            ["Modify markers / detections", "Textura, OBJ con UV, detections_all.csv y detections_mapped_3d.csv."],
            ["Modify labels manually", "OBJ, detections_labeled.csv con etiquetas y XYZ, y un archivo .blend riggeado/etiquetado."],
            ["Modify measures to obtain", "detections_labeled.csv con etiquetas y coordenadas XYZ."],
        ],
        [2.25, 4.25],
    )
    add_callout(
        doc,
        "Compatibilidad.",
        "Si una acción aparece deshabilitada, colocar el cursor sobre el botón puede mostrar qué archivo falta. "
        "Los proyectos recientes incluyen project_context.json para facilitar la reapertura aunque se hayan movido carpetas.",
    )


def add_csv_format(doc: Document) -> None:
    add_heading(doc, "8. Formato de CSV de medidas", 1)
    add_body(
        doc,
        "Los CSV de operaciones definen qué distancias se calculan. Deben estar separados por punto y coma y contener las columnas esperadas. "
        "TYPE y FINGER pueden propagarse desde filas anteriores si hay celdas vacías, pero DIMENSION y los marcadores deben estar definidos.",
    )
    rows = [
        ["TYPE", "Categoría de la medida: LENGTHS, BREADTHS, DEPTHS o EXTRAS."],
        ["FINGER", "Dedo o grupo anatómico asociado a la medida."],
        ["DIMENSION", "Nombre de la medida que aparecerá en los resultados."],
        ["MARKER 1", "Etiqueta del primer punto usado para calcular la distancia."],
        ["MARKER 2", "Etiqueta del segundo punto usado para calcular la distancia."],
    ]
    add_table(doc, ["Columna", "Descripción"], rows, [1.6, 4.9])
    add_callout(
        doc,
        "Coincidencia exacta.",
        "Los valores de MARKER 1 y MARKER 2 deben coincidir exactamente con las etiquetas presentes en detections_labeled.csv.",
    )


def add_faq(doc: Document) -> None:
    add_heading(doc, "9. FAQ y resolución de problemas", 1)
    faqs = [
        (
            "La aplicación no arranca en otro ordenador",
            "Comprobar que se ha copiado la carpeta completa de la release, no sólo HAMLET.exe. "
            "También conviene descomprimir el ZIP en una ruta local con permisos de escritura.",
        ),
        (
            "Blender no se encuentra o aparece acceso denegado",
            "La release debe contener blender/blender.exe. Si se configura una ruta manual, debe apuntar al ejecutable blender.exe o a la carpeta que lo contiene.",
        ),
        (
            "OpenPose no detecta la mano correctamente",
            "Revisar la orientación del modelo, la calidad del render y la presencia de openpose/models/hand/pose_deploy.prototxt y pose_iter_102000.caffemodel.",
        ),
        (
            "Se detectan demasiados falsos positivos",
            "Usar un umbral más restrictivo, reducir tolerancias HSV o seleccionar muestras manuales más limpias y homogéneas.",
        ),
        (
            "Se detectan pocos puntos",
            "Usar tolerancias más amplias, seleccionar más muestras en distintas zonas del marcador y evitar brillos o sombras extremas.",
        ),
        (
            "Un punto aparece en 2D pero no pasa correctamente a 3D",
            "No todos los puntos de la textura pueden proyectarse sobre la malla si el UV o la geometría no lo permiten. Validar siempre la pestaña 3D antes de continuar.",
        ),
        (
            "El etiquetado automático acierta muy poco",
            "La causa más habitual es una orientación incompatible del modelo o un rig incorrecto. Revisar el resultado de OpenPose y corregir etiquetas manualmente.",
        ),
        (
            "No se puede añadir una medida personalizada",
            "Revisar que el nombre no esté vacío, que no exista otra medida con el mismo nombre y que PTO1 y PTO2 sean distintos.",
        ),
        (
            "El CSV de medidas personalizado no carga",
            "Comprobar separador punto y coma, codificación compatible y columnas TYPE, FINGER, DIMENSION, MARKER 1 y MARKER 2.",
        ),
        (
            "El resultado final marca MISSING_MARKER",
            "Una de las etiquetas requeridas por la medida no existe en detections_labeled.csv o quedó vacía tras la corrección manual.",
        ),
    ]
    add_table(doc, ["Problema", "Acción recomendada"], [[q, a] for q, a in faqs], [2.2, 4.3])


def add_contact(doc: Document) -> None:
    add_heading(doc, "10. Contacto", 1)
    add_body(
        doc,
        "Jon Andoni Fernández Moyano\n"
        "Departamento de Biomecánica y Ergonomía, Universitat Jaume I (UJI)\n"
        "moyanoj@uji.es",
    )


def build_docx() -> Path:
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    doc = Document()
    configure_styles(doc)
    add_footer(doc.sections[0], "HAMLET - Guía de usuario")

    add_cover(doc)
    add_index(doc)
    add_summary(doc)
    add_concepts(doc)
    add_requirements(doc)
    add_io(doc)
    add_quick_guide(doc)
    add_extended_guide(doc)
    add_open_project(doc)
    add_csv_format(doc)
    add_faq(doc)
    add_contact(doc)

    doc.save(DOCX_PATH)
    return DOCX_PATH


if __name__ == "__main__":
    print(build_docx())
